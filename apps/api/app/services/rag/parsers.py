"""Belge ayrıştırıcıları.

Desteklenen türler: PDF, DOCX, TXT, Markdown, HTML, XML, JSON, CSV, Python,
JavaScript, TypeScript, SQL, YAML ve log dosyaları.
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

from app.core.logging import get_logger

logger = get_logger(__name__)

TEXT_EXTENSIONS: dict[str, tuple[str, str]] = {
    ".txt": ("documents", "text/plain"),
    ".md": ("documents", "text/markdown"),
    ".markdown": ("documents", "text/markdown"),
    ".log": ("documents", "text/plain"),
    ".json": ("documents", "application/json"),
    ".csv": ("documents", "text/csv"),
    ".yaml": ("documents", "application/yaml"),
    ".yml": ("documents", "application/yaml"),
    ".ini": ("documents", "text/plain"),
    ".env": ("documents", "text/plain"),
    ".py": ("code", "text/x-python"),
    ".js": ("code", "text/javascript"),
    ".jsx": ("code", "text/javascript"),
    ".ts": ("code", "text/typescript"),
    ".tsx": ("code", "text/typescript"),
    ".sql": ("code", "application/sql"),
    ".sh": ("code", "text/x-shellscript"),
    ".ps1": ("code", "text/x-powershell"),
    ".java": ("code", "text/x-java"),
    ".go": ("code", "text/x-go"),
    ".rs": ("code", "text/x-rust"),
    ".c": ("code", "text/x-c"),
    ".cpp": ("code", "text/x-c++"),
    ".h": ("code", "text/x-c"),
    ".css": ("code", "text/css"),
    ".html": ("documents", "text/html"),
    ".htm": ("documents", "text/html"),
    ".toml": ("documents", "text/plain"),
    ".xml": ("documents", "application/xml"),
}

BINARY_EXTENSIONS: dict[str, tuple[str, str]] = {
    ".pdf": ("documents", "application/pdf"),
    ".docx": (
        "documents",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
}

SUPPORTED_EXTENSIONS = {**TEXT_EXTENSIONS, **BINARY_EXTENSIONS}

SKIP_DIRS = {
    "node_modules",
    ".git",
    "__pycache__",
    ".venv",
    "venv",
    "dist",
    "build",
    ".next",
    ".cache",
    "site-packages",
    ".mypy_cache",
    ".ruff_cache",
    ".pytest_cache",
    "release",
    "dist-electron",
}

@dataclass(slots=True)
class ParsedPage:
    """Ayrıştırılmış bir sayfa/bölüm."""

    text: str
    page: int | None = None

@dataclass(slots=True)
class ParsedDocument:
    """Ayrıştırma çıktısı."""

    pages: list[ParsedPage] = field(default_factory=list)
    mime_type: str = "text/plain"
    collection: str = "documents"
    error: str | None = None

    @property
    def full_text(self) -> str:
        """Tüm sayfaların birleşimi."""
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def is_empty(self) -> bool:
        """Kullanılabilir metin var mı?"""
        return not self.full_text.strip()

def is_supported(path: str | Path) -> bool:
    """Uzantı desteklenen türlerden mi?"""
    return Path(path).suffix.lower() in SUPPORTED_EXTENSIONS

def classify(path: str | Path) -> tuple[str, str]:
    """Yol için ``(collection, mime_type)`` döndürür."""
    ext = Path(path).suffix.lower()
    return SUPPORTED_EXTENSIONS.get(ext, ("documents", "application/octet-stream"))

def iter_supported_files(root: Path, recursive: bool = True) -> list[Path]:
    """Klasördeki desteklenen dosyaları listeler."""
    if root.is_file():
        return [root] if is_supported(root) else []

    found: list[Path] = []
    iterator = root.rglob("*") if recursive else root.glob("*")
    for path in iterator:
        if not path.is_file():
            continue
        if any(part in SKIP_DIRS or part.startswith(".") for part in path.parts):
            continue
        if is_supported(path):
            found.append(path)
    return found

async def parse_file(path: Path) -> ParsedDocument:
    """Dosyayı türüne göre ayrıştırır (bloklamayan)."""
    return await asyncio.to_thread(parse_file_sync, path)

def parse_file_sync(path: Path) -> ParsedDocument:
    """Senkron ayrıştırma."""
    ext = path.suffix.lower()
    collection, mime = classify(path)

    try:
        if ext == ".pdf":
            return _parse_pdf(path, collection, mime)
        if ext == ".docx":
            return _parse_docx(path, collection, mime)
        if ext == ".json":
            return _parse_json(path, collection, mime)
        if ext == ".csv":
            return _parse_csv(path, collection, mime)
        if ext in {".html", ".htm"}:
            return _parse_html(path, collection, mime)
        if ext == ".xml":
            return _parse_xml(path, collection, mime)
        return _parse_text(path, collection, mime)
    except Exception as exc:
        logger.warning("parse_failed", path=str(path), error=str(exc))
        return ParsedDocument(mime_type=mime, collection=collection, error=str(exc))

def _read_text(path: Path) -> str:
    """Kodlaması bilinmeyen metin dosyasını okur."""
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    try:
        import chardet

        detected = chardet.detect(raw[:200_000])
        encoding = detected.get("encoding") or "latin-1"
        return raw.decode(encoding, errors="replace")
    except ImportError:
        return raw.decode("latin-1", errors="replace")

def _parse_text(path: Path, collection: str, mime: str) -> ParsedDocument:
    """Düz metin / kod / markdown."""
    text = _read_text(path)
    return ParsedDocument(pages=[ParsedPage(text=text)], mime_type=mime, collection=collection)

_HTML_SKIP = frozenset({"script", "style"})
_HTML_BLOCK = frozenset(
    {
        "p",
        "div",
        "br",
        "tr",
        "li",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "header",
        "footer",
        "nav",
        "table",
        "ul",
        "ol",
        "pre",
        "blockquote",
        "hr",
        "td",
        "th",
    }
)
_HTML_MULTI_NL = re.compile(r"\n{3,}")
_HTML_MULTI_SPACE = re.compile(r"[ \t]{2,}")

_HTML_ATX = {
    "h1": "# ",
    "h2": "## ",
    "h3": "### ",
    "h4": "#### ",
    "h5": "##### ",
    "h6": "###### ",
}

class _HtmlTextParser(HTMLParser):
    """Markitdown: script/style düşür, body varsa onu al, satır satır metin."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip = 0
        self._body_depth = 0
        self._saw_body = False
        self._all: list[str] = []
        self._body: list[str] = []
        self._atx = ""

    def _parts(self) -> list[str]:
        return self._body if self._saw_body else self._all

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _HTML_SKIP:
            self._skip += 1
            return
        if tag == "body":
            self._saw_body = True
            self._body_depth += 1
        if self._skip:
            return
        if tag in _HTML_ATX:
            self._atx = _HTML_ATX[tag]
        if tag in _HTML_BLOCK:
            self._parts().append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _HTML_SKIP:
            self._skip = max(0, self._skip - 1)
            return
        if tag == "body":
            self._body_depth = max(0, self._body_depth - 1)
        if self._skip:
            return
        if tag in _HTML_ATX:
            self._atx = ""
        if tag in _HTML_BLOCK and tag != "br":
            self._parts().append("\n")

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        if tag not in _HTML_SKIP:
            self.handle_endtag(tag)

    def handle_data(self, data: str) -> None:
        if self._skip or (self._saw_body and self._body_depth <= 0):
            return
        text = data.strip()
        if not text:
            return
        if self._atx:
            self._parts().append(self._atx + text)
            self._atx = ""
            return
        self._parts().append(text)

    def text(self) -> str:
        joined = " ".join(self._parts())
        joined = joined.replace("\xa0", " ")
        joined = joined.replace(" \n", "\n").replace("\n ", "\n")
        joined = _HTML_MULTI_SPACE.sub(" ", joined)
        return _HTML_MULTI_NL.sub("\n\n", joined).strip()

def html_to_text(html: str) -> str:
    """HTML'den düz metin — Markitdown ``get_text`` + ATX başlık (stdlib)."""
    if not html.strip():
        return ""
    parser = _HtmlTextParser()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        return html
    return parser.text()

def _parse_html(path: Path, collection: str, mime: str) -> ParsedDocument:
    """HTML — etiket/script düşer, gövde metni kalır."""
    text = html_to_text(_read_text(path))
    return ParsedDocument(pages=[ParsedPage(text=text)], mime_type=mime, collection=collection)

def xml_to_text(raw: str) -> str:
    """XML'den düz metin — Unstructured ``xml_keep_tags=False`` (stdlib).

    LlamaIndex ``ET.tostring`` etiket bırakır; çalınmadı. Harici varlık yok.
    """
    if not raw.strip():
        return ""
    try:
        root = ElementTree.fromstring(raw)
    except ElementTree.ParseError:
        return ""
    parts = [piece.strip() for piece in root.itertext() if piece and piece.strip()]
    return "\n".join(parts)

def _parse_xml(path: Path, collection: str, mime: str) -> ParsedDocument:
    """XML — etiket düşer, yaprak metin kalır."""
    text = xml_to_text(_read_text(path))
    doc = ParsedDocument(pages=[ParsedPage(text=text)], mime_type=mime, collection=collection)
    if doc.is_empty:
        doc.error = "XML'den metin çıkarılamadı."
    return doc

def _parse_pdf(
    path: Path,
    collection: str,
    mime: str,
    *,
    ocr_page: Callable[[object, int], str] | None = None,
) -> ParsedDocument:
    """PDF — metin; boş sayfada RapidOCR (görüntü). Tesseract yok."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[ParsedPage] = []
    for number, page in enumerate(reader.pages, start=1):
        text = pdf_page_text(page, number, ocr_page=ocr_page)
        if text.strip():
            pages.append(ParsedPage(text=text, page=number))

    doc = ParsedDocument(pages=pages, mime_type=mime, collection=collection)
    if doc.is_empty:
        doc.error = (
            "PDF'ten metin çıkarılamadı. Belge taranmış görüntü olabilir ve OCR sonuç vermedi."
        )
    return doc

def pdf_page_text(
    page: object,
    number: int,
    *,
    ocr_page: Callable[[object, int], str] | None = None,
) -> str:
    """Sayfa metni; yoksa OCR (testte enjekte edilebilir)."""
    try:
        text = page.extract_text() or ""  # type: ignore[attr-defined]
    except Exception:
        text = ""
    if text.strip():
        return text
    if ocr_page is not None:
        try:
            return ocr_page(page, number) or ""
        except Exception as exc:
            logger.warning("pdf_ocr_hook_failed", page=number, error=str(exc))
            return ""
    return _ocr_pdf_page(page)

def _page_image_bytes(page: object) -> list[bytes]:
    blobs: list[bytes] = []
    try:
        images = page.images  # type: ignore[attr-defined]
    except Exception:
        return blobs
    for image in images or []:
        data = getattr(image, "data", None)
        if data:
            blobs.append(bytes(data))
    return blobs

_ocr_engine: object | None = None

def _ocr_pdf_page(page: object) -> str:
    """Sayfadaki gömülü görüntülere RapidOCR. Model yoksa boş döner."""
    parts = [_rapidocr_bytes(blob) for blob in _page_image_bytes(page)]
    return "\n".join(part for part in parts if part.strip())

def _rapidocr_result_text(result: object) -> str:
    if result is None:
        return ""
    txts = getattr(result, "txts", None)
    if txts:
        return "\n".join(str(item) for item in txts if item)
    if isinstance(result, tuple) and result:
        result = result[0]
    if isinstance(result, list):
        lines: list[str] = []
        for item in result:
            if isinstance(item, (list, tuple)) and len(item) >= 2:
                lines.append(str(item[1]))
            elif isinstance(item, str):
                lines.append(item)
        return "\n".join(lines)
    return str(result).strip()

def _rapidocr_bytes(image_bytes: bytes) -> str:
    global _ocr_engine
    try:
        import numpy as np
        from PIL import Image
        from rapidocr import RapidOCR
    except ImportError:
        logger.warning("rapidocr_missing")
        return ""
    try:
        if _ocr_engine is None:
            _ocr_engine = RapidOCR()
        image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
        result = _ocr_engine(np.array(image))
        return _rapidocr_result_text(result)
    except Exception as exc:
        logger.warning("rapidocr_failed", error=str(exc))
        return ""

def _parse_docx(path: Path, collection: str, mime: str) -> ParsedDocument:
    """DOCX — paragraf ve tablo metni."""
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return ParsedDocument(
        pages=[ParsedPage(text="\n".join(parts))], mime_type=mime, collection=collection
    )

def _parse_json(path: Path, collection: str, mime: str) -> ParsedDocument:
    """JSON — okunabilir biçimde düzleştirir."""
    text = _read_text(path)
    try:
        data = json.loads(text)
        pretty = json.dumps(data, ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        pretty = text
    return ParsedDocument(pages=[ParsedPage(text=pretty)], mime_type=mime, collection=collection)

def _parse_csv(path: Path, collection: str, mime: str) -> ParsedDocument:
    """CSV — satırları ``sütun: değer`` biçimine çevirir."""
    text = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel  # type: ignore[assignment]

    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        return ParsedDocument(mime_type=mime, collection=collection)

    header = rows[0]
    lines: list[str] = [" | ".join(header)]
    for row in rows[1:]:
        pairs = [f"{h}: {v}" for h, v in zip(header, row, strict=False) if v]
        if pairs:
            lines.append("; ".join(pairs))
    return ParsedDocument(
        pages=[ParsedPage(text="\n".join(lines))], mime_type=mime, collection=collection
    )
